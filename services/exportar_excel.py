"""Genera el archivo de Excel con todas las hojas (Fase 13).

Usa openpyxl. Cada hoja lleva encabezados en negrita con color, columnas
ajustadas, filtros, la primera fila inmovilizada y formato de dinero/fecha.
"""

from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from services.fabricados import calcular_fabricado
from services.recetas import calcular_receta

COLOR_ENCABEZADO = "4F46E5"


def _a_fecha(texto):
    if not texto:
        return None
    try:
        return datetime.strptime(str(texto)[:10], "%Y-%m-%d").date()
    except ValueError:
        return str(texto)


def _a_fecha_hora(texto):
    if not texto:
        return None
    for formato in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(texto), formato)
        except ValueError:
            pass
    return str(texto)


def _agregar_hoja(wb, titulo, encabezados, filas, formatos=None):
    ws = wb.create_sheet(titulo)
    ws.append(encabezados)
    for fila in filas:
        ws.append(fila)

    for columna in range(1, len(encabezados) + 1):
        celda = ws.cell(row=1, column=columna)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=COLOR_ENCABEZADO)
        celda.alignment = Alignment(horizontal="center")

    if formatos:
        for indice, formato in enumerate(formatos, start=1):
            if not formato:
                continue
            for fila in range(2, ws.max_row + 1):
                celda = ws.cell(row=fila, column=indice)
                if formato == "money":
                    celda.number_format = '"$"#,##0.00'
                elif formato == "date":
                    celda.number_format = "YYYY-MM-DD"
                elif formato == "datetime":
                    celda.number_format = "YYYY-MM-DD HH:MM"

    for columna in range(1, len(encabezados) + 1):
        letra = get_column_letter(columna)
        ancho = len(str(encabezados[columna - 1]))
        for fila in range(2, ws.max_row + 1):
            valor = ws.cell(row=fila, column=columna).value
            if valor is not None:
                ancho = max(ancho, len(str(valor)))
        ws.column_dimensions[letra].width = min(ancho + 2, 45)

    ws.freeze_panes = "A2"
    if ws.max_row >= 2:
        ws.auto_filter.ref = "A1:%s%d" % (
            get_column_letter(len(encabezados)), ws.max_row)
    return ws


def _hoja_proveedores(con, wb):
    filas = con.execute(
        "SELECT id, nombre, razon_social, cuit_rut, telefono, whatsapp, email,"
        " direccion, persona_contacto, notas, creado_en FROM proveedores ORDER BY nombre"
    ).fetchall()
    datos = [[f["id"], f["nombre"], f["razon_social"] or "", f["cuit_rut"] or "",
              f["telefono"] or "", f["whatsapp"] or "", f["email"] or "",
              f["direccion"] or "", f["persona_contacto"] or "", f["notas"] or "",
              _a_fecha_hora(f["creado_en"])] for f in filas]
    _agregar_hoja(wb, "Proveedores",
                  ["ID", "Nombre", "Razón social", "CUIT/RUT", "Teléfono", "WhatsApp",
                   "Email", "Dirección", "Contacto", "Notas", "Creado"],
                  datos, [None] * 10 + ["datetime"])


def _hoja_productos(con, wb):
    filas = con.execute(
        "SELECT p.id, p.nombre, COALESCE(c.nombre, '') AS categoria,"
        " uc.abreviatura AS ucompra, uu.abreviatura AS uso, p.precio_actual,"
        " COALESCE(pr.nombre, '') AS proveedor, COALESCE(p.sku, '') AS sku,"
        " COALESCE(p.notas, '') AS notas, p.creado_en"
        " FROM productos p"
        " LEFT JOIN categorias c ON c.id = p.categoria_id"
        " JOIN unidades uc ON uc.id = p.unidad_compra_id"
        " JOIN unidades uu ON uu.id = p.unidad_uso_id"
        " LEFT JOIN proveedores pr ON pr.id = p.proveedor_principal_id"
        " ORDER BY p.nombre").fetchall()
    datos = [[f["id"], f["nombre"], f["categoria"], f["ucompra"], f["uso"],
              f["precio_actual"], f["proveedor"], f["sku"], f["notas"],
              _a_fecha_hora(f["creado_en"])] for f in filas]
    _agregar_hoja(wb, "Productos",
                  ["ID", "Nombre", "Categoría", "Unidad compra", "Unidad uso",
                   "Precio actual", "Proveedor", "SKU", "Notas", "Creado"],
                  datos, [None, None, None, None, None, "money", None, None, None, "datetime"])


def _hoja_compras(con, wb):
    filas = con.execute(
        "SELECT c.id, pr.nombre AS proveedor, c.fecha, COALESCE(c.numero_documento, ''),"
        " COALESCE(c.observaciones, ''), c.total"
        " FROM compras c JOIN proveedores pr ON pr.id = c.proveedor_id"
        " ORDER BY c.fecha DESC, c.id DESC").fetchall()
    datos = [[f["id"], f["proveedor"], _a_fecha(f["fecha"]), f[3], f[4], f["total"]]
             for f in filas]
    _agregar_hoja(wb, "Compras",
                  ["ID", "Proveedor", "Fecha", "Nº factura/remito", "Observaciones", "Total"],
                  datos, [None, None, "date", None, None, "money"])


def _hoja_detalle_compras(con, wb):
    filas = con.execute(
        "SELECT d.id, d.compra_id, p.nombre AS producto, d.cantidad, u.abreviatura,"
        " d.precio_total, d.precio_unitario FROM detalle_compras d"
        " JOIN productos p ON p.id = d.producto_id"
        " JOIN unidades u ON u.id = d.unidad_id ORDER BY d.compra_id, d.id").fetchall()
    datos = [[f["id"], f["compra_id"], f["producto"], f["cantidad"], f["abreviatura"],
              f["precio_total"], f["precio_unitario"]] for f in filas]
    _agregar_hoja(wb, "Detalle Compras",
                  ["ID", "Compra", "Producto", "Cantidad", "Unidad", "Precio total",
                   "Precio unitario"],
                  datos, [None, None, None, None, None, "money", "money"])


def _hoja_historial(con, wb):
    filas = con.execute(
        "SELECT h.id, p.nombre AS producto, h.fecha, h.precio_unitario, u.abreviatura,"
        " h.detalle_compra_id FROM historial_precios h"
        " JOIN productos p ON p.id = h.producto_id"
        " JOIN unidades u ON u.id = h.unidad_id"
        " ORDER BY h.fecha DESC, h.id DESC").fetchall()
    datos = [[f["id"], f["producto"], _a_fecha(f["fecha"]), f["precio_unitario"],
              f["abreviatura"], f["detalle_compra_id"]] for f in filas]
    _agregar_hoja(wb, "Historial Precios",
                  ["ID", "Producto", "Fecha", "Precio unitario", "Unidad", "Detalle compra"],
                  datos, [None, None, "date", "money", None, None])


def _hoja_recetas(con, wb):
    unidades = {u["id"]: u for u in con.execute("SELECT * FROM unidades")}
    datos = []
    for receta in con.execute("SELECT * FROM recetas ORDER BY nombre"):
        calculo = calcular_receta(con, receta["id"], unidades=unidades)
        datos.append([receta["id"], receta["nombre"], receta["descripcion"] or "",
                      receta["rendimiento_cantidad"], receta["rendimiento_unidad"],
                      receta["notas"] or "", float(calculo["total"]),
                      float(calculo["costo_unidad"])])
    _agregar_hoja(wb, "Recetas",
                  ["ID", "Nombre", "Descripción", "Rendimiento", "Unidad",
                   "Notas", "Costo total", "Costo por unidad"],
                  datos, [None, None, None, None, None, None, "money", "money"])

    ingredientes = con.execute(
        "SELECT r.nombre AS receta, p.nombre AS producto, ri.cantidad, u.abreviatura"
        " FROM receta_ingredientes ri"
        " JOIN recetas r ON r.id = ri.receta_id"
        " JOIN productos p ON p.id = ri.producto_id"
        " JOIN unidades u ON u.id = ri.unidad_id ORDER BY r.nombre, ri.id").fetchall()
    datos_ing = [[f["receta"], f["producto"], f["cantidad"], f["abreviatura"]]
                 for f in ingredientes]
    _agregar_hoja(wb, "Ingredientes",
                  ["Receta", "Producto", "Cantidad", "Unidad"], datos_ing)


def _hoja_fabricados(con, wb):
    datos = []
    precios = []
    for fila in con.execute("SELECT id FROM productos_fabricados ORDER BY nombre"):
        calculo = calcular_fabricado(con, fila["id"])
        datos.append([calculo["id"], calculo["nombre"], calculo["receta"],
                      calculo["cantidad_fabricada"], float(calculo["costo_lote"]),
                      float(calculo["costo_unidad"]), calculo["notas"] or ""])
        for presentacion in calculo["presentaciones"]:
            precios.append([
                calculo["nombre"], presentacion["nombre"],
                presentacion["cantidad_unidades"], float(presentacion["costo"]),
                float(presentacion["tienda"]) if presentacion["tienda"] is not None else None,
                float(presentacion["delivery"]) if presentacion["delivery"] is not None else None,
                float(presentacion["ganancia"]) if presentacion["ganancia"] is not None else None,
            ])

    _agregar_hoja(wb, "Productos Fabricados",
                  ["ID", "Nombre", "Receta", "Cantidad fabricada", "Costo lote",
                   "Costo por unidad", "Notas"],
                  datos, [None, None, None, None, "money", "money", None])
    _agregar_hoja(wb, "Costos y Precios",
                  ["Producto", "Presentación", "Unidades", "Costo", "🏪 Tienda",
                   "🛵 Delivery", "Ganancia"],
                  precios, [None, None, None, "money", "money", "money", "money"])


def construir_libro(con):
    wb = Workbook()
    wb.remove(wb.active)
    _hoja_proveedores(con, wb)
    _hoja_productos(con, wb)
    _hoja_compras(con, wb)
    _hoja_detalle_compras(con, wb)
    _hoja_historial(con, wb)
    _hoja_recetas(con, wb)
    _hoja_fabricados(con, wb)
    return wb


# FIN services/exportar_excel.py
